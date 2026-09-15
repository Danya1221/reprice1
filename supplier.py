"""Collect fresh multipart replies, including edits and supported price documents."""
import asyncio
import io
import re
import time
from contextlib import suppress

from telethon import events, functions, utils
from telethon.errors import BotResponseTimeoutError
from telethon.tl import types

from prices import CLOSED, parse_documents, clean


class SupplierTimeout(RuntimeError):
    pass


def signature(message):
    buttons = tuple(tuple(b.text for b in row) for row in (message.buttons or []))
    document = getattr(message, "document", None)
    return message.raw_text or "", buttons, getattr(document, "id", None)


def button_label(text):
    return clean(re.sub(r"[^\w\s]", " ", text)).casefold()


def find_button(message, wanted):
    wanted = button_label(wanted)
    for row in message.buttons or []:
        for button in row:
            if button_label(button.text or "") == wanted:
                return button
    return None


def _username(value):
    if not isinstance(value, str):
        return None
    text = value.strip()
    match = re.fullmatch(r"(?:https?://)?t\.me/([A-Za-z0-9_]+)/?", text, re.I)
    if match:
        text = match.group(1)
    text = text.lstrip("@")
    if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{3,}", text):
        return text
    return None


def _resolved_entity(result):
    peer = result.peer
    if isinstance(peer, types.PeerUser):
        return next((item for item in result.users if item.id == peer.user_id), None)
    if isinstance(peer, types.PeerChannel):
        return next((item for item in result.chats if item.id == peer.channel_id), None)
    if isinstance(peer, types.PeerChat):
        return next((item for item in result.chats if item.id == peer.chat_id), None)
    return None


async def resolve_input_peer(client, value, *, label="Telegram peer", dialog_limit=400):
    """Resolve a configured peer without depending on Telethon's transient entity cache.

    Public usernames are resolved explicitly with ResolveUsernameRequest, which returns
    the access_hash required for bots/users/channels. Numeric/private peers fall back to
    the current dialog list only when the direct lookup cannot work from StringSession.
    """
    username = _username(value)
    if username:
        try:
            result = await client(functions.contacts.ResolveUsernameRequest(username))
            entity = _resolved_entity(result)
            if entity is None:
                raise RuntimeError("Telegram вернул username без полной entity")
            return utils.get_input_peer(entity)
        except Exception as exc:
            raise RuntimeError(
                f"{label}: не удалось открыть @{username}: {type(exc).__name__}: {exc}"
            ) from exc

    try:
        return await client.get_input_entity(value)
    except Exception as direct_error:
        wanted_id = None
        try:
            wanted_id = utils.get_peer_id(value)
        except Exception:
            pass

        try:
            async for dialog in client.iter_dialogs(limit=dialog_limit):
                entity = dialog.entity
                if wanted_id is not None and utils.get_peer_id(entity) == wanted_id:
                    return utils.get_input_peer(entity)
        except Exception as dialog_error:
            raise RuntimeError(
                f"{label}: не удалось найти {value!r} в Telegram: "
                f"{type(dialog_error).__name__}: {dialog_error}"
            ) from dialog_error

        raise RuntimeError(
            f"{label}: Telegram знает ID {value!r}, но не дал access_hash. "
            "Укажи публичный @username или открой этот чат на аккаунте поставщиков. "
            f"Исходная ошибка: {type(direct_error).__name__}: {direct_error}"
        ) from direct_error


class SupplierReader:
    def __init__(self, client, settings, source):
        self.client = client
        self.settings = settings
        self.source = source
        self.entity = None

    async def resolve(self):
        self.entity = await resolve_input_peer(
            self.client,
            self.source.peer,
            label=self.source.label,
        )

    def is_closed_text(self, text):
        custom = clean(self.source.closed_text).casefold()
        return bool(CLOSED.search(text or "") or (custom and custom in clean(text or "").casefold()))

    async def collect_action(self, action):
        """Snapshot before action, listen before sending, poll as a fallback."""
        s = self.settings
        before = await self.client.get_messages(self.entity, limit=s.history_limit)
        baseline = {m.id: signature(m) for m in before if not m.out}
        max_id = max((m.id for m in before), default=0)
        collected = {}
        changed_at = None

        def accept(message):
            nonlocal changed_at
            if message.out:
                return
            sig = signature(message)
            if message.id in baseline and baseline[message.id] == sig:
                return
            if message.id <= max_id and message.id not in baseline:
                return
            previous = collected.get(message.id)
            if previous is None or signature(previous) != sig:
                collected[message.id] = message
                changed_at = time.monotonic()

        async def on_message(event):
            accept(event.message)

        new_event = events.NewMessage(chats=self.entity, incoming=True)
        edit_event = events.MessageEdited(chats=self.entity, incoming=True)
        self.client.add_event_handler(on_message, new_event)
        self.client.add_event_handler(on_message, edit_event)
        try:
            try:
                await asyncio.wait_for(action(), timeout=s.response_timeout)
            except BotResponseTimeoutError:
                # A callback may deliver a new message but omit answerCallbackQuery.
                pass
            deadline = time.monotonic() + s.response_timeout
            while time.monotonic() < deadline:
                recent = await self.client.get_messages(self.entity, limit=s.history_limit)
                for message in recent:
                    if message.id > max_id or message.id in baseline:
                        accept(message)
                if collected and changed_at is not None and time.monotonic() - changed_at >= s.quiet_seconds:
                    return [collected[k] for k in sorted(collected)]
                await asyncio.sleep(s.action_delay)
            if collected:
                raise SupplierTimeout("Ответ поставщика не завершён: увеличь RESPONSE_TIMEOUT")
            raise SupplierTimeout("Поставщик не прислал нового ответа; старое меню не будет опубликовано")
        finally:
            self.client.remove_event_handler(on_message, new_event)
            self.client.remove_event_handler(on_message, edit_event)

    async def messages(self):
        if self.entity is None:
            await self.resolve()
        if self.source.mode == "feed":
            messages = [m for m in await self.client.get_messages(
                self.entity, limit=self.settings.history_limit)]
            if not messages:
                raise SupplierTimeout("Лента поставщика пуста")
            latest = next((m for m in messages if (m.raw_text or "").strip() or m.document), None)
            if latest and self.is_closed_text(latest.raw_text):
                return [latest]
            return list(reversed(messages))
        if not self.source.request:
            raise ValueError("Для SOURCE_MODE=bot нужен REQUEST_TEXT")
        messages = await self.collect_action(
            lambda: self.client.send_message(self.entity, self.source.request, parse_mode=None))
        for wanted in self.source.buttons:
            if any(self.is_closed_text(m.raw_text) for m in messages):
                return messages
            button = next((b for m in reversed(messages) if (b := find_button(m, wanted))), None)
            if button is None:
                labels = [b.text for m in messages for row in (m.buttons or []) for b in row]
                raise RuntimeError(f"Нет кнопки «{wanted}». Доступны: {', '.join(labels)}")
            if not isinstance(button.button, (types.KeyboardButton, types.KeyboardButtonCallback)):
                raise RuntimeError(f"Кнопка «{wanted}» не является текстовой или callback-кнопкой")
            messages = await self.collect_action(button.click)
        return messages

    async def document_text(self, message):
        document = getattr(message, "document", None)
        if not document:
            return ""
        if document.size > 5 * 1024 * 1024:
            raise RuntimeError("Файл прайса больше 5 МБ")
        name = next((a.file_name for a in document.attributes if hasattr(a, "file_name")), "").lower()
        if not name.endswith((".txt", ".csv", ".xlsx")):
            raise RuntimeError("Поддерживаются текстовые прайсы, TXT, CSV и XLSX; получен другой файл")
        payload = await self.client.download_media(message, file=bytes)
        if not payload:
            raise RuntimeError("Не удалось скачать файл прайса")
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            import zipfile
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    raise RuntimeError("Распакованный XLSX превышает 50 МБ")
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            try:
                lines = []
                for sheet in workbook:
                    for row in sheet.iter_rows(values_only=True):
                        cells = [str(value).strip() for value in row if value is not None]
                        if cells:
                            lines.append(" ".join(cells))
                        if len(lines) > 30000:
                            raise RuntimeError("В XLSX больше 30000 строк")
                return "\n".join(lines)
            finally:
                workbook.close()
        text = None
        for encoding in ("utf-8-sig", "cp1251"):
            with suppress(UnicodeDecodeError):
                text = payload.decode(encoding)
                break
        if text is None:
            raise RuntimeError("Не удалось прочитать кодировку TXT/CSV")
        if name.endswith(".csv"):
            import csv
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
                rows = csv.reader(io.StringIO(text), dialect)
                text = "\n".join(" ".join(cell.strip() for cell in row if cell.strip()) for row in rows)
            except csv.Error:
                pass
        return text

    async def fetch(self):
        messages = await self.messages()
        documents = []
        for message in messages:
            if message.raw_text:
                documents.append(message.raw_text)
            if message.document:
                documents.append(await self.document_text(message))
        result = parse_documents(documents, self.settings.currency)
        if not result.items and any(self.is_closed_text(text) for text in documents):
            result.closed = True
        if not result.items and not result.closed:
            raise RuntimeError("В ответе нет распознанных товаров. Прайс в канале сохранён")
        return result
